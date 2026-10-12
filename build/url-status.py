"""Dental refocus URL report: every URL from the pre-refocus sitemap plus the new pages, with its
final status (kept / 301 to / new / draft), checked against the live site.

Usage: python build/url-status.py <old-urls.txt> [https://rcm.texassolutions.co]
Writes docs/dental-refocus-url-status.xlsx and .csv.
"""
import csv
import json
import sys
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / 'wordpress' / 'rcm-core' / 'data'
old_file = sys.argv[1]
base = (sys.argv[2] if len(sys.argv) > 2 else 'https://rcm.texassolutions.co').rstrip('/')


def load(name):
    return json.loads((DATA / f'{name}.json').read_text(encoding='utf-8'))


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *a, **k):
        return None


opener = urllib.request.build_opener(NoRedirect)


def check(path):
    req = urllib.request.Request(base + path, headers={'User-Agent': 'RCM-url-status/1.0'})
    try:
        r = opener.open(req, timeout=60)
        return r.status, ''
    except urllib.error.HTTPError as e:
        return e.code, e.headers.get('Location') or ''
    except Exception as e:  # network error
        return 0, str(e)


redirects = {r['from']: r['to'] for r in load('redirects')}
labels = {}
for s in load('services'):
    labels[f"/services/{s['slug']}/"] = ('Service', s.get('name', s['slug']))
for s in load('specialties'):
    labels[f"/specialties/{s['slug']}/"] = ('Specialty', s.get('name', s['slug']))
for t in load('tools'):
    labels[f"/tools/{t['slug']}/"] = ('Tool', t.get('name', t['slug']))
for b in load('blog'):
    labels[f"/blog/{b['slug']}/"] = ('Blog post', b.get('title', b['slug']))
drafts = set()
for p in load('pages'):
    path = '/' + (p['parent'] + '/' if p.get('parent') else '') + p['key'] + '/'
    labels[path] = ('Page', p.get('title', p['key']))
    if p.get('status') == 'draft':
        drafts.add(path)
for c in ['dental-billing', 'denial-management', 'practice-management', 'compliance-and-regulation']:
    labels[f'/blog/category/{c}/'] = ('Blog category', c.replace('-', ' ').title())
labels['/'] = ('Home', 'Home')

old = [urllib.request.urlparse(u.strip()).path for u in Path(old_file).read_text().split() if u.strip()]
old_set = set(old)
new_paths = [p for p in labels if p not in old_set and p not in drafts]

rows = []
for path in sorted(old_set) + sorted(new_paths) + sorted(drafts - old_set):
    in_old = path in old_set
    if path in redirects:
        status, target = '301 redirect', redirects[path]
    elif path in drafts:
        status, target = 'Draft (not published)', ''
    elif in_old:
        status, target = 'Kept (same URL)', ''
    else:
        status, target = 'New page', ''
    code, loc = check(path)
    if status == '301 redirect':
        ok = code == 301 and loc == base + target
        tcode, _ = check(target)
        ok = ok and tcode == 200
        live = f'{code} -> {tcode}'
    elif status.startswith('Draft'):
        ok, live = code == 404, str(code)
    else:
        ok, live = code == 200, str(code)
    kind, name = labels.get(path, labels.get(target, ('', '')))
    rows.append([base + path, 'Yes' if in_old else 'No', status, base + target if target else '', kind, name, live, 'OK' if ok else 'CHECK'])
    print(path, status, live, 'OK' if ok else 'CHECK')

head = ['URL', 'In old sitemap', 'Final status', 'Redirects to', 'Type', 'Name (kept or target)', 'Live HTTP check', 'Result']
out = ROOT / 'docs'
with open(out / 'dental-refocus-url-status.csv', 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f)
    w.writerow(head)
    w.writerows(rows)

wb = Workbook()
ws = wb.active
ws.title = 'URL status'
ws.append(head)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='0F1F3F')
for r in rows:
    ws.append(r)
    if r[-1] != 'OK':
        ws.cell(ws.max_row, len(head)).fill = PatternFill('solid', fgColor='FDE68A')
for i, wdt in enumerate([62, 14, 22, 62, 14, 48, 16, 10], 1):
    ws.column_dimensions[get_column_letter(i)].width = wdt
ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions

s = wb.create_sheet('Summary')
counts = {}
for r in rows:
    counts[r[2]] = counts.get(r[2], 0) + 1
s.append(['Final status', 'URLs'])
for k, v in sorted(counts.items()):
    s.append([k, v])
s.append(['Checks not OK', sum(1 for r in rows if r[-1] != 'OK')])
s.column_dimensions['A'].width = 28
wb.save(out / 'dental-refocus-url-status.xlsx')
print(counts, 'not ok:', sum(1 for r in rows if r[-1] != 'OK'))
